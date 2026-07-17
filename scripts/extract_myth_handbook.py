"""Convert the editorial handbook's 60 single-cell cards to CMS-ready JSON."""

import json
import re
import sys
from pathlib import Path
from docx import Document

LABELS = [
    "MYTH",
    "EVIDENCE-BASED REALITY",
    "EXPLANATION",
    "INDIA CONTEXT",
    "EVIDENCE TO VERIFY",
    "EDITORIAL CAUTION",
]

def parse_card(text: str, index: int) -> dict:
    pattern = "(" + "|".join(re.escape(label) for label in LABELS) + ")"
    parts = re.split(pattern, text.strip())
    values: dict[str, str] = {}
    for pos in range(1, len(parts), 2):
        values[parts[pos]] = " ".join(parts[pos + 1].split())
    caution = values.get("EDITORIAL CAUTION", "")
    caution = re.sub(r"Last reviewed:.*$", "", caution).strip()
    return {
        "id": f"myth-{index:02d}",
        "myth": values.get("MYTH", ""),
        "reality": values.get("EVIDENCE-BASED REALITY", ""),
        "explanation": values.get("EXPLANATION", ""),
        "indiaContext": values.get("INDIA CONTEXT", ""),
        "evidence": values.get("EVIDENCE TO VERIFY", ""),
        "caution": caution,
        "reviewed": "July 2026",
        "status": "Editorial research draft",
    }

source, target = map(Path, sys.argv[1:3])
document = Document(source)
cards = [parse_card(table.cell(0, 0).text, index) for index, table in enumerate(document.tables, 1)]
if len(cards) != 60 or any(not card["myth"] or not card["reality"] for card in cards):
    raise SystemExit("Handbook structure changed: expected 60 complete cards")
target.write_text(json.dumps(cards, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(f"Wrote {len(cards)} cards to {target}")
